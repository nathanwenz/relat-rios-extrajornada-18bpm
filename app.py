import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import io

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def gerar_relatorio_word(df_escala, data_extenso="23 de setembro de 2026 (quarta-feira)"):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Cabeçalho Institucional
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_hdr.add_run("ESTADO DO PARANÁ POLÍCIA MILITAR\n2º COMANDO REGIONAL DE POLÍCIA MILITAR\n18° BATALHÃO DE POLÍCIA MILITAR\n")
    r1.bold = True
    
    r2 = p_hdr.add_run("PROGRAMAÇÃO EXTRAJORNADA VOLUNTÁRIA\n")
    r2.bold = True
    r2.font.color.rgb = RGBColor(0, 32, 96)

    p_data = doc.add_paragraph()
    p_data.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_data.add_run(data_extenso).bold = True

    p_obs = doc.add_paragraph()
    p_obs.add_run("* As equipes ficarão à disposição do CPU, Adjunto e COPOM para atendimentos de ocorrências. Na ausência de ocorrências, deverão seguir o cartão de programa.").italic = True

    # Agrupamento por VOUCHER e CIDADE
    for (voucher, cidade), grupo in df_escala.groupby(["VOUCHER", "CIDADE"]):
        primeiro = grupo.iloc
        data_str = str(primeiro.get("DATA", ""))
        hora_ini = str(primeiro.get("HORA", "18:00"))
        hora_fim = str(primeiro.get("HORA_FIM", "23:59"))
        
        table = doc.add_table(rows=7 + len(grupo), cols=2)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        campos = [
            ("CIDADE/VOUCHER", f"{cidade} - VOUCHER {voucher}"),
            ("DATA", data_str),
            ("LOCAL", cidade),
            ("HORÁRIO", f"Das {hora_ini} às {hora_fim}")
        ]

        for i, (label, val) in enumerate(campos):
            r = table.rows[i]
            r.cells.paragraphs.add_run(label).bold = True
            r.cells[2].paragraphs.add_run(val)
            set_cell_background(r.cells, "F2F2F2")

        # Texto fixo padrão
        r4 = table.rows
        c_m = r4.cells
        c_m.merge(r4.cells[2])
        c_m.paragraphs.add_run("A equipe ficará a Disposição do COPOM e CPU ou Adjunto. | SISGCOP 61076\nEquipe deverá fazer contato com o Adjunto ao assumir serviço. A equipe além de realizar o atendimento de ocorrências, deverá realizar o Patrulhamento Ostensivo e Preventivo na área designada para atuar.")
        set_cell_background(c_m, "FAFAFA")

        # Lista de Policiais
        r5 = table.rows
        c_pm = r5.cells
        c_pm.merge(r5.cells[2])
        c_pm.paragraphs.add_run("POLICIAIS MILITARES ESCALADOS (PM VOLUNTÁRIOS)").bold = True
        set_cell_background(c_pm, "D9E1F2")

        for idx, (_, row) in enumerate(grupo.iterrows()):
            r_pm = table.rows[6 + idx]
            r_pm.cells.paragraphs.add_run(str(row["Posto/Grad."])).bold = True
            r_pm.cells[2].paragraphs.add_run(f"{row['PM VOLUNTÁRIO']} — CPF: {row['CPF']}")

        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Interface Web Streamlit
st.title("🛡️ Gerador de Relatório Extrajornada PM")
arquivo = st.file_uploader("Envie a planilha de escala (.xlsx ou .csv)", type=["xlsx", "csv"])

if arquivo:
    df = pd.read_excel(arquivo) if arquivo.name.endswith(".xlsx") else pd.read_csv(arquivo)
    data_cabecalho = st.text_input("Data para o cabeçalho", "23 de setembro de 2026 (quarta-feira)")
    
    if st.button("Gerar Documento Word"):
        docx_bytes = gerar_relatorio_word(df, data_cabecalho)
        st.download_button("📥 Baixar Relatório Preenchido (.docx)", docx_bytes, "RELATORIO_EXTRAJORNADA.docx")
